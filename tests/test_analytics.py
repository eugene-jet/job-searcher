"""Unit tests for the time-series analytics (CSV source of truth + workbook)."""

import openpyxl

import analytics


# --- _as_int ---------------------------------------------------------------

def test_as_int_parses_and_treats_blanks_as_gaps():
    assert analytics._as_int("5") == 5
    assert analytics._as_int("0") == 0
    assert analytics._as_int("") is None
    assert analytics._as_int("x") is None
    assert analytics._as_int(None) is None


# --- record_day ------------------------------------------------------------

def _csv(tmp_path):
    return str(tmp_path / "vacancy_counts.csv")


def test_record_day_writes_sorted_rows(tmp_path):
    csv_path = _csv(tmp_path)
    analytics.record_day({"dou": 200, "djinni": 80}, "2026-09-11", csv_path)
    analytics.record_day({"dou": 210, "djinni": 90}, "2026-09-10", csv_path)

    rows = analytics._read_rows(csv_path)
    assert rows == {
        "2026-09-10": {"dou": "210", "djinni": "90"},
        "2026-09-11": {"dou": "200", "djinni": "80"},
    }
    # Dates are written in ascending order regardless of insertion order.
    with open(csv_path, encoding="utf-8") as fh:
        body = fh.read()
    assert body.index("2026-09-10") < body.index("2026-09-11")


def test_record_day_upserts_same_day(tmp_path):
    csv_path = _csv(tmp_path)
    analytics.record_day({"dou": 200, "djinni": 80}, "2026-09-11", csv_path)
    analytics.record_day({"dou": 205, "djinni": 85}, "2026-09-11", csv_path)
    rows = analytics._read_rows(csv_path)
    assert rows == {"2026-09-11": {"dou": "205", "djinni": "85"}}


def test_record_day_none_does_not_overwrite_existing_value(tmp_path):
    csv_path = _csv(tmp_path)
    analytics.record_day({"dou": 200, "djinni": 80}, "2026-09-11", csv_path)
    # A failed DOU scrape (None) must leave the previously recorded 200 intact.
    analytics.record_day({"dou": None, "djinni": 81}, "2026-09-11", csv_path)
    rows = analytics._read_rows(csv_path)
    assert rows == {"2026-09-11": {"dou": "200", "djinni": "81"}}


def test_record_day_none_without_prior_value_is_a_blank_gap(tmp_path):
    csv_path = _csv(tmp_path)
    analytics.record_day({"dou": None, "djinni": 80}, "2026-09-11", csv_path)
    rows = analytics._read_rows(csv_path)
    assert rows == {"2026-09-11": {"dou": "", "djinni": "80"}}


# --- build_workbook --------------------------------------------------------

def test_build_workbook_writes_table_and_chart(tmp_path):
    csv_path = _csv(tmp_path)
    xlsx_path = str(tmp_path / "analytics.xlsx")
    analytics.record_day({"dou": 200, "djinni": 80}, "2026-09-10", csv_path)
    analytics.record_day({"dou": None, "djinni": 81}, "2026-09-11", csv_path)

    analytics.build_workbook(csv_path, xlsx_path)

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb["Counts"]
    assert [c.value for c in ws[1]] == ["date", "DOU", "Djinni"]
    assert [c.value for c in ws[2]] == ["2026-09-10", 200, 80]
    # The blank DOU cell becomes None (a chart gap), not a real zero.
    assert [c.value for c in ws[3]] == ["2026-09-11", None, 81]
    assert len(ws._charts) == 1


def test_build_workbook_is_safe_when_csv_missing(tmp_path):
    xlsx_path = str(tmp_path / "empty.xlsx")
    analytics.build_workbook(str(tmp_path / "does-not-exist.csv"), xlsx_path)
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb["Counts"]
    assert [c.value for c in ws[1]] == ["date", "DOU", "Djinni"]
    assert ws.max_row == 1  # header only, no data and no chart
    assert len(ws._charts) == 0
